from io import BytesIO
import math
import os
from zipfile import BadZipFile

from flask import Blueprint, abort, flash, redirect, request, send_file, url_for
from flask_login import current_user
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.utils import secure_filename

from extensions import db
import models

spreadsheets_bp = Blueprint("spreadsheets", __name__)
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _add_sheet(workbook, title, headers, rows):
    sheet = workbook.active if workbook.active.title == "Sheet" else workbook.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append([
            "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            for value in row
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _workbook_response(workbook, filename):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype=XLSX_MIMETYPE,
        as_attachment=True,
        download_name=filename,
    )


@spreadsheets_bp.route("/exports/<table_name>")
def export_table(table_name):
    workspace_id = current_user.workspace_id
    workbook = Workbook()

    if table_name == "ingredients":
        ingredients = models.Ingredient.query.filter_by(
            workspace_id=workspace_id
        ).order_by(models.Ingredient.name.asc()).all()
        _add_sheet(workbook, "Ingredients", [
            "Name", "Category", "Vendor", "Package Cost", "Quantity per Package",
            "Unit Cost", "Photo URL",
        ], [
            [
                item.name,
                item.type_info.name if item.type_info else "",
                item.vendor.name if item.vendor else "",
                item.cost_per_pkg,
                item.qty_per_pkg,
                item.cost_per_pkg / item.qty_per_pkg if item.qty_per_pkg else 0,
                url_for("ingredients.ingredient_image", ingredient_id=item.id)
                if item.image_filename else "",
            ]
            for item in ingredients
        ])
    elif table_name in {"products", "recipes"}:
        products = models.Product.query.filter_by(
            workspace_id=workspace_id
        ).order_by(models.Product.name.asc()).all()
        _add_sheet(workbook, "Products", [
            "Product", "Code", "Design Type", "Raw Wholesale Cost",
            "Marked-Up Materials", "Design Fee", "Suggested Retail Price",
            "Actual Price", "Margin %",
        ], [
            [
                product.name,
                product.product_code,
                product.design_style.name if product.design_style else "",
                product.raw_wholesale_cost,
                product.marked_up_cost,
                product.design_fee,
                product.suggested_retail_price,
                product.actual_price,
                product.profit_margin_percentage,
            ]
            for product in products
        ])
        _add_sheet(workbook, "Recipes", [
            "Product", "Product Code", "Ingredient", "Quantity", "Photo URL",
        ], [
            [
                product.name,
                product.product_code,
                line.ingredient.name,
                line.quantity,
                url_for(
                    "ingredients.ingredient_image",
                    ingredient_id=line.ingredient.id,
                ) if line.ingredient.image_filename else "",
            ]
            for product in products
            for line in product.sorted_recipe_items
        ])
    elif table_name == "vendors":
        vendors = models.Vendor.query.filter_by(
            workspace_id=workspace_id
        ).order_by(models.Vendor.name.asc()).all()
        _add_sheet(workbook, "Vendors", ["Vendor", "Ingredient Count"], [
            [vendor.name, len(vendor.ingredients)] for vendor in vendors
        ])
    elif table_name == "ingredient-types":
        types = models.IngredientType.query.filter_by(
            workspace_id=workspace_id
        ).order_by(models.IngredientType.name.asc()).all()
        _add_sheet(workbook, "Ingredient Types", ["Name", "Markup Multiplier"], [
            [ingredient_type.name, ingredient_type.markup] for ingredient_type in types
        ])
    elif table_name == "design-types":
        types = models.DesignType.query.filter_by(
            workspace_id=workspace_id
        ).order_by(models.DesignType.name.asc()).all()
        _add_sheet(workbook, "Design Fees", ["Name", "Design Fee Percentage"], [
            [design_type.name, design_type.design_fee_percentage] for design_type in types
        ])
    elif table_name == "margins":
        products = models.Product.query.filter_by(
            workspace_id=workspace_id
        ).order_by(models.Product.name.asc()).all()
        _add_sheet(workbook, "Margins", [
            "Product", "Code", "Raw Wholesale Cost", "Design Fee",
            "Suggested Retail Price", "Margin %",
        ], [
            [
                product.name,
                product.product_code,
                product.raw_wholesale_cost,
                product.design_fee,
                product.suggested_retail_price,
                product.profit_margin_percentage,
            ]
            for product in products
        ])
    else:
        abort(404)

    return _workbook_response(workbook, f"{table_name}.xlsx")


@spreadsheets_bp.route("/exports/vendors/<int:vendor_id>")
def export_vendor_report(vendor_id):
    vendor = models.Vendor.query.filter_by(
        id=vendor_id, workspace_id=current_user.workspace_id
    ).first_or_404()
    ingredients = models.Ingredient.query.filter_by(
        vendor_id=vendor.id, workspace_id=current_user.workspace_id
    ).order_by(models.Ingredient.name.asc()).all()
    workbook = Workbook()
    _add_sheet(workbook, "Vendor Inventory", [
        "Category", "Ingredient", "Package Quantity", "Package Cost", "Photo URL",
    ], [
        [
            item.type_info.name if item.type_info else "",
            item.name,
            item.qty_per_pkg,
            item.cost_per_pkg,
            url_for("ingredients.ingredient_image", ingredient_id=item.id)
            if item.image_filename else "",
        ]
        for item in ingredients
    ])
    vendor_filename = secure_filename(vendor.name) or "vendor"
    return _workbook_response(
        workbook, f"{vendor_filename}-inventory.xlsx"
    )


@spreadsheets_bp.route("/imports/ingredients", methods=["POST"])
def import_ingredients():
    uploaded_file = request.files.get("file")
    if uploaded_file is None or not uploaded_file.filename:
        abort(400, description="Choose an .xlsx file to import.")
    if os.path.splitext(uploaded_file.filename)[1].lower() != ".xlsx":
        abort(400, description="Ingredient imports must be .xlsx files.")

    contents = uploaded_file.stream.read(10 * 1024 * 1024 + 1)
    if len(contents) > 10 * 1024 * 1024:
        abort(413, description="Excel files must be 10 MB or smaller.")
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, ValueError):
        abort(400, description="The selected file is not a valid Excel workbook.")

    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = {
        str(value).strip().casefold(): index
        for index, value in enumerate(header_row)
        if value is not None
    }
    name_column = headers.get("name")
    type_column = headers.get("category", headers.get("ingredient type"))
    cost_column = headers.get("package cost", headers.get("cost per package"))
    quantity_column = headers.get(
        "quantity per package", headers.get("qty per package")
    )
    vendor_column = headers.get("vendor")
    if None in {name_column, type_column, cost_column, quantity_column}:
        workbook.close()
        abort(
            400,
            description=(
                "The first row must contain Name, Category, Package Cost, and "
                "Quantity per Package columns. Download the Ingredients export "
                "for a compatible template."
            ),
        )

    workspace_id = current_user.workspace_id
    known_names = {
        name.casefold()
        for (name,) in db.session.query(models.Ingredient.name).filter_by(
            workspace_id=workspace_id
        ).all()
    }
    type_ids = {
        item.name.casefold(): item.id
        for item in models.IngredientType.query.filter_by(workspace_id=workspace_id).all()
    }
    vendor_ids = {
        item.name.casefold(): item.id
        for item in models.Vendor.query.filter_by(workspace_id=workspace_id).all()
    }
    pending = []
    row_errors = []
    duplicate_count = 0

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(value is None or str(value).strip() == "" for value in row):
            continue
        try:
            if row[name_column] is None or not str(row[name_column]).strip():
                raise ValueError("name is required")
            name = str(row[name_column]).strip()
            if len(name) > 100:
                raise ValueError("name must be at most 100 characters")
            normalized_name = name.casefold()
            if normalized_name in known_names:
                duplicate_count += 1
                continue
            type_name = str(row[type_column]).strip()
            package_cost = float(row[cost_column])
            quantity = float(row[quantity_column])
            vendor_name = (
                str(row[vendor_column]).strip()
                if vendor_column is not None and row[vendor_column] is not None
                else ""
            )
            if not math.isfinite(package_cost) or package_cost < 0:
                raise ValueError("package cost must be a finite number at least zero")
            if not math.isfinite(quantity) or quantity <= 0:
                raise ValueError("quantity per package must be greater than zero")
            ingredient_type_id = type_ids.get(type_name.casefold())
            if ingredient_type_id is None:
                raise ValueError(f"ingredient type '{type_name}' is not in this workspace")
            vendor_id = vendor_ids.get(vendor_name.casefold()) if vendor_name else None
            if vendor_name and vendor_id is None:
                raise ValueError(f"vendor '{vendor_name}' is not in this workspace")
            known_names.add(normalized_name)
            pending.append(models.Ingredient(
                workspace_id=workspace_id,
                name=name,
                cost_per_pkg=package_cost,
                qty_per_pkg=quantity,
                ingredient_type_id=ingredient_type_id,
                vendor_id=vendor_id,
            ))
        except (IndexError, TypeError, ValueError) as error:
            row_errors.append(f"Row {row_number}: {error}")

    workbook.close()
    if row_errors:
        details = "; ".join(row_errors[:5])
        if len(row_errors) > 5:
            details += f"; and {len(row_errors) - 5} more row(s)"
        flash(f"No ingredients were imported. Fix these Excel errors: {details}", "error")
    else:
        db.session.add_all(pending)
        db.session.commit()
        flash(
            f"Imported {len(pending)} ingredient(s); skipped {duplicate_count} existing item(s).",
            "success",
        )
    return redirect(url_for("ingredients.ingredients"))
