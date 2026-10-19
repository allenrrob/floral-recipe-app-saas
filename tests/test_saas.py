import os
import re
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
os.environ.setdefault("SECRET_KEY", "test-only-secret")

from app import create_app
from extensions import db
import models


class SaaSTestCase(unittest.TestCase):
    def setUp(self):
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test-only-secret",
                "SQLALCHEMY_DATABASE_URI": "sqlite://",
                "WTF_CSRF_ENABLED": False,
                "STRIPE_SECRET_KEY": "sk_test",
                "STRIPE_PRICE_ID": "price_test",
                "STRIPE_WEBHOOK_SECRET": "whsec_test",
            }
        )
        self.client = self.app.test_client()
        with self.app.app_context():
            db.create_all()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()

    def create_account(self, client, email, workspace):
        response = client.post(
            "/signup",
            data={
                "workspace_name": workspace,
                "email": email,
                "password": "a-long-test-password",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/billing/status")
        with self.app.app_context():
            user = models.User.query.filter_by(email=email).one()
            subscription = user.workspace.subscription
            self.assertEqual(subscription.status, "incomplete")
            subscription.status = "active"
            db.session.commit()
            return user.workspace_id

    def test_signup_creates_independent_workspace_and_login(self):
        workspace_id = self.create_account(self.client, "owner@example.test", "Petal Studio")
        with self.app.app_context():
            user = models.User.query.filter_by(email="owner@example.test").one()
            self.assertEqual(user.workspace_id, workspace_id)
            self.assertNotEqual(user.password_hash, "a-long-test-password")
            self.assertEqual(
                models.IngredientType.query.filter_by(workspace_id=workspace_id).count(),
                4,
            )
            type_id = models.IngredientType.query.filter_by(
                workspace_id=workspace_id
            ).first().id
        self.client.post(
            "/ingredients",
            data={
                "action": "create_ingredient",
                "name": "Test Rose",
                "cost_per_pkg": "32",
                "qty_per_pkg": "25",
                "ingredient_type_id": str(type_id),
                "vendor_id": "",
            },
        )

        response = self.client.get("/ingredients")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"ingredient-edit-form", response.data)
        css = self.client.get("/static/css/main.css")
        self.assertIn(b"grid-column-end: -1", css.get_data())
        css.close()

        self.client.post("/logout")
        response = self.client.post(
            "/login",
            data={"email": "owner@example.test", "password": "a-long-test-password"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/billing/status")

    def test_unpaid_workspace_is_redirected_to_billing(self):
        response = self.client.post(
            "/signup",
            data={
                "workspace_name": "New Florist",
                "email": "new@example.test",
                "password": "a-long-test-password",
            },
        )
        self.assertEqual(response.status_code, 302)
        response = self.client.get("/")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/billing/status")

    def test_mutating_requests_require_csrf_tokens(self):
        self.app.config["WTF_CSRF_ENABLED"] = True
        response = self.client.post(
            "/signup",
            data={
                "workspace_name": "Missing Token Florist",
                "email": "missing-token@example.test",
                "password": "a-long-test-password",
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_workspace_records_cannot_be_read_or_edited_by_another_workspace(self):
        first_workspace = self.create_account(
            self.client, "first@example.test", "First Florist"
        )
        with self.app.app_context():
            first_type_id = models.IngredientType.query.filter_by(
                workspace_id=first_workspace
            ).first().id
        self.client.post("/vendors", data={"action": "create_vendor", "name": "First Vendor"})
        with self.app.app_context():
            first_vendor = models.Vendor.query.filter_by(
                workspace_id=first_workspace
            ).one()
            first_vendor_id = first_vendor.id

        self.client.post("/logout")
        second_client = self.app.test_client()
        second_workspace = self.create_account(
            second_client, "second@example.test", "Second Florist"
        )
        self.assertNotEqual(first_workspace, second_workspace)

        response = second_client.get(f"/vendors/{first_vendor_id}/inventory")
        self.assertEqual(response.status_code, 404)

        second_client.post(
            "/vendors",
            data={
                "action": "edit_vendor",
                "vendor_id": str(first_vendor_id),
                "name": "Hijacked Vendor",
            },
        )
        response = second_client.post(
            "/ingredients",
            data={
                "action": "create_ingredient",
                "name": "Foreign Rose",
                "cost_per_pkg": "20",
                "qty_per_pkg": "10",
                "ingredient_type_id": str(first_type_id),
                "vendor_id": "",
            },
        )
        self.assertEqual(response.status_code, 404)

        with self.app.app_context():
            first_vendor = db.session.get(models.Vendor, first_vendor_id)
            self.assertEqual(first_vendor.name, "First Vendor")
            self.assertEqual(
                models.Ingredient.query.filter_by(workspace_id=second_workspace).count(),
                0,
            )

    def test_invited_user_only_sees_recipes_and_suggested_price(self):
        workspace_id = self.create_account(
            self.client, "owner@example.test", "Floral Studio"
        )
        with self.app.app_context():
            ingredient_type = models.IngredientType.query.filter_by(
                workspace_id=workspace_id, name="Standard Flower"
            ).one()
            design_type = models.DesignType.query.filter_by(
                workspace_id=workspace_id
            ).first()
            ingredient = models.Ingredient(
                workspace_id=workspace_id,
                name="Confidential Rose",
                cost_per_pkg=20,
                qty_per_pkg=10,
                ingredient_type_id=ingredient_type.id,
            )
            product = models.Product(
                workspace_id=workspace_id,
                name="Private Arrangement",
                product_code="PRV-01",
                design_type_id=design_type.id,
            )
            db.session.add_all([ingredient, product])
            db.session.flush()
            db.session.add(models.ProductIngredient(
                workspace_id=workspace_id,
                product_id=product.id,
                ingredient_id=ingredient.id,
                quantity=3,
            ))
            ingredient_id = ingredient.id
            design_type_id = design_type.id
            db.session.commit()

        invitation_response = self.client.post(
            "/team",
            data={
                "action": "invite",
                "email": "designer@example.test",
                "role": "user",
            },
        )
        self.assertEqual(invitation_response.status_code, 200)
        invite_url = re.search(
            r'value="(http[^"]+)"', invitation_response.get_data(as_text=True)
        ).group(1)
        invite_path = urlsplit(invite_url).path
        invited_client = self.app.test_client()
        response = invited_client.post(
            invite_path, data={"password": "a-different-long-password"}
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/recipes")
        self.assertEqual(invited_client.get("/login").location, "/recipes")
        self.assertEqual(invited_client.get("/signup").location, "/recipes")
        self.assertEqual(invited_client.get("/").location, "/recipes")

        response = invited_client.get("/recipes")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Confidential Rose", response.data)
        self.assertNotIn(b"Raw Wholesale Cost", response.data)
        self.assertNotIn(b"Suggested Retail Price", response.data)
        self.assertEqual(invited_client.get("/ingredients").status_code, 403)
        self.assertEqual(invited_client.get("/exports/ingredients").status_code, 403)
        self.assertEqual(invited_client.get("/quick-calculator").status_code, 403)

        calculator = invited_client.get("/designer-calculator")
        self.assertEqual(calculator.status_code, 200)
        self.assertNotIn(b"$2.00/unit", calculator.data)
        self.assertNotIn(b"Labor Fee", calculator.data)
        result = invited_client.post(
            "/designer-calculator",
            data={
                "action": "calculate",
                "row_count": "1",
                "design_type_id": str(design_type_id),
                "ingredient_id[]": str(ingredient_id),
                "quantity[]": "3",
            },
        )
        self.assertEqual(result.status_code, 200)
        self.assertIn(b"$25.20", result.data)
        self.assertNotIn(b"Raw Wholesale Cost", result.data)
        self.assertNotIn(b"Marked-Up Materials", result.data)
        self.assertNotIn(b"Design Fee", result.data)

        with self.app.app_context():
            designer_id = models.User.query.filter_by(
                email="designer@example.test"
            ).one().id
        self.client.post(
            "/team",
            data={
                "action": "change_role",
                "user_id": str(designer_id),
                "role": "admin",
            },
        )
        self.assertEqual(invited_client.get("/ingredients").status_code, 200)

    def test_ingredient_photos_excel_import_export_and_margin_metrics(self):
        workspace_id = self.create_account(
            self.client, "excel@example.test", "Excel Florist"
        )
        with self.app.app_context():
            ingredient_type = models.IngredientType.query.filter_by(
                workspace_id=workspace_id, name="Standard Flower"
            ).one()
            ingredient_type_id = ingredient_type.id

        response = self.client.post(
            "/ingredients",
            data={
                "action": "create_ingredient",
                "name": "Test Rose",
                "cost_per_pkg": "32",
                "qty_per_pkg": "25",
                "ingredient_type_id": str(ingredient_type_id),
                "vendor_id": "",
                "image": (BytesIO(b"\x89PNG\r\n\x1a\nmock-image"), "rose.png"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        with self.app.app_context():
            ingredient = models.Ingredient.query.filter_by(
                workspace_id=workspace_id, name="Test Rose"
            ).one()
            ingredient_id = ingredient.id
            self.assertTrue(ingredient.image_filename)
            image_path = (
                Path(self.app.instance_path)
                / "ingredient-uploads"
                / str(workspace_id)
                / ingredient.image_filename
            )
        image_response = self.client.get(f"/ingredient-images/{ingredient_id}")
        self.assertEqual(image_response.status_code, 200)
        self.assertEqual(image_response.data, b"\x89PNG\r\n\x1a\nmock-image")
        image_response.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Category", "Vendor", "Package Cost", "Quantity per Package"])
        sheet.append(["Test Rose", "Not a real type", "", 100, 10])
        sheet.append(["Imported Orchid", "Standard Flower", "", 25, 10])
        upload = BytesIO()
        workbook.save(upload)
        upload.seek(0)
        response = self.client.post(
            "/imports/ingredients",
            data={"file": (upload, "ingredients.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        with self.app.app_context():
            imported = models.Ingredient.query.filter_by(
                workspace_id=workspace_id, name="Imported Orchid"
            ).one()
            self.assertEqual(imported.cost_per_pkg, 25)
            self.assertEqual(
                models.Ingredient.query.filter_by(
                    workspace_id=workspace_id, name="Test Rose"
                ).count(),
                1,
            )
            vendor = models.Vendor(name="Photo Vendor", workspace_id=workspace_id)
            db.session.add(vendor)
            db.session.flush()
            existing_ingredient = models.Ingredient.query.filter_by(
                workspace_id=workspace_id, name="Test Rose"
            ).one()
            existing_ingredient.vendor_id = vendor.id
            vendor_id = vendor.id
            db.session.commit()

        export_response = self.client.get("/exports/ingredients")
        self.assertEqual(export_response.status_code, 200)
        exported = load_workbook(BytesIO(export_response.data), read_only=True)
        self.assertEqual(exported.active["A1"].value, "Name")
        self.assertEqual(exported.active.max_row, 3)
        export_response.close()
        for table_name in (
            "products", "vendors", "ingredient-types", "design-types", "margins"
        ):
            response = self.client.get(f"/exports/{table_name}")
            self.assertEqual(response.status_code, 200, table_name)
            response.close()

        report = self.client.get(f"/vendors/{vendor_id}/inventory")
        self.assertIn(b"Test Rose", report.data)
        self.assertIn(f"/ingredient-images/{ingredient_id}".encode(), report.data)
        vendor_export = self.client.get(f"/exports/vendors/{vendor_id}")
        self.assertEqual(vendor_export.status_code, 200)
        self.assertTrue(vendor_export.data.startswith(b"PK"))
        vendor_export.close()

        margins = self.client.get("/profit-margins")
        self.assertEqual(margins.status_code, 200)
        self.assertNotIn(b"Total Wholesale Value", margins.data)
        self.assertNotIn(b"Total Retail Value", margins.data)
        image_path.unlink(missing_ok=True)

    def test_signed_stripe_webhook_updates_subscription_idempotently(self):
        workspace_id = self.create_account(
            self.client, "billing@example.test", "Billing Florist"
        )
        event = {
            "id": "evt_subscription_change",
            "created": 100,
            "type": "customer.subscription.updated",
            "data": {
                "object": {
                    "id": "sub_123",
                    "customer": "cus_123",
                    "status": "active",
                    "metadata": {"workspace_id": str(workspace_id)},
                }
            },
        }
        with patch("routes.billing.stripe.Webhook.construct_event", return_value=event):
            response = self.client.post(
                "/billing/webhook",
                data=b"signed-event",
                headers={"Stripe-Signature": "test-signature"},
            )
        self.assertEqual(response.status_code, 200)

        duplicate = {
            **event,
            "data": {
                "object": {
                    **event["data"]["object"],
                    "status": "canceled",
                }
            },
        }
        with patch("routes.billing.stripe.Webhook.construct_event", return_value=duplicate):
            response = self.client.post(
                "/billing/webhook",
                data=b"same-event",
                headers={"Stripe-Signature": "test-signature"},
            )
        self.assertEqual(response.status_code, 200)
        stale_event = {
            **event,
            "id": "evt_stale_subscription_change",
            "created": 99,
            "data": {
                "object": {
                    **event["data"]["object"],
                    "status": "canceled",
                }
            },
        }
        with patch("routes.billing.stripe.Webhook.construct_event", return_value=stale_event):
            response = self.client.post(
                "/billing/webhook",
                data=b"stale-event",
                headers={"Stripe-Signature": "test-signature"},
            )
        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            subscription = models.Subscription.query.filter_by(
                workspace_id=workspace_id
            ).one()
            self.assertEqual(subscription.status, "active")
            self.assertEqual(subscription.stripe_subscription_id, "sub_123")


if __name__ == "__main__":
    unittest.main()
